import io
import json
import re
from types import SimpleNamespace

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pipeline import telco_state

from app import sessions
from app.agents.report import report_node
from app.exports.excel import SHEETS, write_excel
from app.exports.pdf import write_pdf
from app.graph.state import ChurnState
from app.main import create_app

REC = {"id": "R1", "priority": 1, "group": "quick_win", "action": "Offer annual plans",
       "problem": "p", "target_segment": "Month-to-month", "effort": "low",
       "customers_affected": {"source_key": "k", "value": 3861, "display": "3,861"},
       "impact": {"source_key": "i", "value": 153.7, "assumption": "10% fewer churners"}}
INSIGHT = {"id": "I1", "title": "Month-to-month churns most", "text": "39.81% churn.",
           "figures": [], "significant": True, "causality_note": "Association only."}


def _values():
    return {**telco_state(), "final_recommendations": [REC], "final_insights": [INSIGHT],
            "validation_report": {"checked": 2, "passed": 2, "failed": 0, "dropped": 0}}


def pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page(?![a-zA-Z])", pdf))


def test_excel_sheets_rows_and_formatting(tmp_path):
    values = _values()
    path = write_excel(values, tmp_path / "r.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == list(SHEETS)
    clean = pd.read_parquet(values["clean_path"])
    preds = pd.read_parquet(values["predictions_path"])
    tests = values["hypothesis_results"]["tests"]
    expected_rows = {"Cleaned_Data": len(clean), "Predictions": len(preds),
                     "Hypothesis_Tests": len(tests),
                     "Drivers": len(values["feature_importance"]["driver_impact"]),
                     "Recommendations": 1}
    for name, rows in expected_rows.items():
        ws = wb[name]
        assert ws.max_row - 1 == rows, name
        assert ws.freeze_panes == "A2"
        assert len(ws.tables) == 1 and not ws.merged_cells.ranges
        assert all(c.font.bold for c in ws[1])
    headers = [c.value for c in wb["Predictions"][1]]
    col = headers.index("churn_probability") + 1
    assert wb["Predictions"].cell(row=2, column=col).number_format == "0.00%"
    col = [c.value for c in wb["Hypothesis_Tests"][1]].index("p_adjusted") + 1
    assert wb["Hypothesis_Tests"].cell(row=2, column=col).number_format == "0.0000"
    assert wb["Recommendations"]["C2"].value == "Offer annual plans"


def test_excel_neutralises_formula_text(tmp_path):
    values = {**_values(), "final_recommendations": [{**REC, "action": "=HYPERLINK(\"x\")"}]}
    wb = load_workbook(write_excel(values, tmp_path / "r.xlsx"))
    assert wb["Recommendations"]["C2"].value == "'=HYPERLINK(\"x\")"


def test_pdf_is_non_empty_with_expected_pages(tmp_path):
    data = write_pdf(_values(), tmp_path / "r.pdf").read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 20_000
    assert 4 <= pages(data) <= 8


def test_report_node_marks_exports_ready():
    state = ChurnState(**{k: v for k, v in telco_state().items() if k in ChurnState.model_fields})
    out = report_node(state)
    assert out["report_paths"]["formats"] == ["excel", "pdf"]
    assert report_node(ChurnState(session_id="x"))["progress"][0].status == "skipped"


class FakeGraph:
    def __init__(self, store):
        self.store = store

    def get_state(self, config):
        values, pending = self.store.get(config["configurable"]["thread_id"], ({}, ()))
        return SimpleNamespace(values=values, next=pending,
                               config={"configurable": {"checkpoint_id": "cp-1"}})


@pytest.fixture
def api():
    app = create_app()
    store = {}
    app.state.runs = SimpleNamespace(graph=FakeGraph(store))
    with TestClient(app) as client:
        yield client, store


def _session(store, values, pending=()):
    sid = sessions.new_session_id()
    folder = sessions.session_dir(sid)
    folder.mkdir(parents=True)
    (folder / sessions.META_FILE).write_text(json.dumps({}))
    store[sid] = (values, pending)
    return sid


def test_export_endpoints(api):
    client, store = api
    sid = _session(store, _values())
    res = client.get(f"/export/{sid}/excel")
    assert res.status_code == 200
    assert "churnlens_report.xlsx" in res.headers["content-disposition"]
    assert load_workbook(io.BytesIO(res.content)).sheetnames == list(SHEETS)
    cached = sessions.session_dir(sid) / "export_cp-1.xlsx"
    stamp = cached.stat().st_mtime_ns
    assert client.get(f"/export/{sid}/excel").status_code == 200
    assert cached.stat().st_mtime_ns == stamp  # built once, then served from the folder
    pdf = client.get(f"/export/{sid}/pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


def test_export_errors(api):
    client, store = api
    assert client.get(f"/export/{'e' * 32}/pdf").status_code == 410
    running = _session(store, _values(), pending=("validator",))
    assert client.get(f"/export/{running}/excel").status_code == 409
    empty = _session(store, {"session_id": "x"})
    assert client.get(f"/export/{empty}/pdf").status_code == 409
    assert client.get(f"/export/{running}/csv").status_code == 404

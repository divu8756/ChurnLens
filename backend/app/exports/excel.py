"""Power BI-ready Excel workbook: one Excel table per sheet, a header row, frozen panes,
number formats, no merged cells (runbook T7.2)."""

import re
import warnings
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

SHEETS = ("Cleaned_Data", "Predictions", "Hypothesis_Tests", "Drivers", "Recommendations")
PERCENT_HINT = re.compile(r"(rate|share|pct_of|lift_share|probability)$", re.I)
P_VALUE_HINT = re.compile(r"^(p_value|p_adjusted)$")
MAX_WIDTH = 60


def _cell(value: Any) -> Any:
    """Excel-safe scalar; text that starts like a formula is written as text."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().replace(tzinfo=None)
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    if isinstance(value, list | dict):
        return str(value)
    return value


def _number_format(column: str, series: pd.Series) -> str | None:
    if not pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series):
        return None
    if P_VALUE_HINT.match(column):
        return "0.0000"
    if PERCENT_HINT.search(column):
        return "0.00%"
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    return "#,##0.00"


def _sheet(wb: Workbook, name: str, frame: pd.DataFrame) -> None:
    """Stream one sheet (openpyxl write-only mode keeps memory flat for large files)."""
    ws = wb.create_sheet(name)
    frame = frame.reset_index(drop=True)
    if not len(frame.columns):
        frame = pd.DataFrame({"note": ["No rows for this session."]})
    columns = [str(c) for c in frame.columns]
    # Write-only sheets need widths and panes before the first row.
    ws.freeze_panes = "A2"
    formats = []
    for i, col in enumerate(columns, start=1):
        formats.append(_number_format(col, frame[col]))
        sample = frame[col].astype(str).head(200)
        width = max([len(col), *sample.str.len().tolist()]) if len(sample) else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(MAX_WIDTH, width + 2)

    def styled(value: Any, fmt: str | None = None, bold: bool = False) -> WriteOnlyCell:
        cell = WriteOnlyCell(ws, value=value)
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = Font(bold=True)
        return cell

    ws.append([styled(c, bold=True) for c in columns])
    for row in frame.itertuples(index=False, name=None):
        ws.append([styled(_cell(v), fmt) if fmt else _cell(v)
                   for v, fmt in zip(row, formats, strict=True)])
    last = f"{get_column_letter(len(columns))}{max(2, len(frame) + 1)}"
    table = Table(displayName=f"t_{name}", ref=f"A1:{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table.tableColumns = [TableColumn(id=i, name=c) for i, c in enumerate(columns, start=1)]
    with warnings.catch_warnings():
        # openpyxl always warns in write-only mode; the columns are set just above.
        warnings.filterwarnings("ignore", message="In write-only mode you must add table columns")
        ws.add_table(table)


def _unique_headers(frame: pd.DataFrame) -> pd.DataFrame:
    """Excel tables need unique, non-empty header names."""
    seen: dict[str, int] = {}
    names = []
    for col in frame.columns:
        base = str(col).strip() or "column"
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return frame.set_axis(names, axis=1)


def workbook_frames(values: dict[str, Any]) -> dict[str, pd.DataFrame]:
    clean = pd.read_parquet(values["clean_path"]) if values.get("clean_path") else pd.DataFrame()
    preds = (pd.read_parquet(values["predictions_path"]) if values.get("predictions_path")
             and Path(values["predictions_path"]).exists() else pd.DataFrame())
    tests = pd.DataFrame([{
        "variable": t.get("variable"), "test": t.get("test_name"),
        "statistic_name": t.get("statistic_name"), "statistic": t.get("statistic"),
        "df": t.get("df"), "p_value": t.get("p_value"), "p_adjusted": t.get("p_adjusted"),
        "significant": t.get("significant"),
        "effect_size_name": (t.get("effect_size") or {}).get("name"),
        "effect_size": (t.get("effect_size") or {}).get("value"),
        "effect_band": (t.get("effect_size") or {}).get("band"),
    } for t in (values.get("hypothesis_results") or {}).get("tests") or []])
    drivers = pd.DataFrame((values.get("feature_importance") or {}).get("driver_impact") or [])
    recs = pd.DataFrame([{
        "priority": r.get("priority"), "group": r.get("group"), "action": r.get("action"),
        "problem": r.get("problem"), "target_segment": r.get("target_segment"),
        "customers_affected": (r.get("customers_affected") or {}).get("value"),
        "impact": (r.get("impact") or {}).get("value"),
        "impact_assumption": (r.get("impact") or {}).get("assumption"),
        "effort": r.get("effort"),
    } for r in values.get("final_recommendations") or []])
    frames = {"Cleaned_Data": clean, "Predictions": preds, "Hypothesis_Tests": tests,
              "Drivers": drivers, "Recommendations": recs}
    return {name: _unique_headers(frame) for name, frame in frames.items()}


def write_excel(values: dict[str, Any], path: Path) -> Path:
    wb = Workbook(write_only=True)
    for name, frame in workbook_frames(values).items():
        _sheet(wb, name, frame)
    wb.save(path)
    return path
